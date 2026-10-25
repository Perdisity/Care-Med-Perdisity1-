import json
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import openpyxl
import pytest

from extensions import db
from models.product import (
    CustomerDeposit,
    Expense,
    InventoryLog,
    Payment,
    Product,
    Rental,
    TankStatus,
    Transaction,
)
from models.settings import SystemSetting
from models.users import SecurityLog


def login(client, email, password):
    return client.post(
        "/login",
        data={"email": email, "password": password},
        follow_redirects=True,
    )


def logout(client):
    return client.get("/logout", follow_redirects=True)


def create_delivery_rental(
    client,
    app_instance,
    product_id,
    customer_id,
    distance,
    quantity=1,
    start_date="2026-10-01",
    return_date="2026-11-30",
):
    response = client.post(
        "/admin/process-rental",
        data={
            "product_ids": str(product_id),
            f"quantity_{product_id}": str(quantity),
            f"unit_price_{product_id}": "700",
            "start_date": start_date,
            "return_date": return_date,
            "duration_unit": "months",
            "customer_id": str(customer_id),
            "fulfillment_type": "Delivery",
            "distance_km": str(distance),
            "delivery_address": "Jaro, Iloilo City",
            "delivery_fee": "100",
            "customer_deposit": "0",
            "amount_paid": "0",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )
    with app_instance.app_context():
        txn = Transaction.query.order_by(Transaction.id.desc()).first()
        rental = txn.rentals[0]
        return response, txn.id, rental.id, Decimal(str(txn.deposit_amount)), txn.deposit_rule


def test_distance_based_deposit_tiers(client, app_instance, seed_data):
    """FR02, RULE-03, TC-03: Distance tiers calculate correct deposit."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    tiers = [
        ("2", Decimal("500.00"), "Up to 3 km"),
        ("4.5", Decimal("1000.00"), "3"),
        ("8", Decimal("1500.00"), "Above 6 km"),
    ]
    for distance, expected_deposit, rule_fragment in tiers:
        _, txn_id, rental_id, deposit, rule = create_delivery_rental(
            client, app_instance, seed_data["product_id"], seed_data["customer_id"], distance
        )
        assert deposit == expected_deposit
        assert rule is not None and rule_fragment in rule


def test_staff_cannot_process_damaged_return(client, app_instance, seed_data):
    """RULE-10, TC-11: Staff users cannot assess damaged returns."""
    with app_instance.app_context():
        txn = Transaction.query.filter(Transaction.deposit_amount == Decimal("500.00")).order_by(Transaction.id.asc()).first()
        rental_id = txn.rentals[0].id
        txn_id = txn.id

    logout(client)
    login(client, "test.staff@gmail.com", "Staff@12345")

    client.post(
        f"/admin/transaction/{txn_id}/return",
        data={
            "returned_items": str(rental_id),
            f"qty_{rental_id}": "1",
            f"status_{rental_id}": "Full",
            f"condition_{rental_id}": "Major Damage",
            f"damage_reason_{rental_id}": "Cracked",
            "return_notes": "Cracked casing",
            "late_fees": "0",
            "refund_deposit": "yes",
            "deposit_refund_amount": "500.00",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        rental = db.session.get(Rental, rental_id)
        assert rental.status == "Active"
        assert Decimal(str(rental.damage_fee or 0)) == Decimal("0.00")
        assert rental.damage_approved_at is None


def test_admin_damaged_return_with_reason(client, app_instance, seed_data):
    """RULE-10, TC-11: Administrator can assess damage with reason code and fee capped at deposit."""
    with app_instance.app_context():
        txn = Transaction.query.filter(Transaction.deposit_amount == Decimal("500.00")).order_by(Transaction.id.asc()).first()
        rental_id = txn.rentals[0].id
        txn_id = txn.id

    logout(client)
    login(client, "test.admin@gmail.com", "Admin@12345")

    client.post(
        f"/admin/transaction/{txn_id}/return",
        data={
            "returned_items": str(rental_id),
            f"qty_{rental_id}": "1",
            f"status_{rental_id}": "Full",
            f"condition_{rental_id}": "Major Damage",
            f"damage_reason_{rental_id}": "Cracked",
            "return_notes": "Cracked casing",
            "late_fees": "0",
            "refund_deposit": "yes",
            "deposit_refund_amount": "500.00",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        rental = db.session.get(Rental, rental_id)
        assert rental.status == "Returned"
        assert rental.damage_reason == "Cracked"
        assert Decimal(str(rental.damage_fee)) == Decimal("250.00")
        assert rental.damage_approved_by_id == seed_data["admin_id"]
        assert rental.damage_approved_at is not None

        approved_log = SecurityLog.query.filter_by(event_type="Damage Approved").first()
        assert approved_log is not None
        assert "Cracked" in (approved_log.description or "")

        refund = Payment.query.filter_by(transaction_id=txn_id, payment_type="Deposit Refund").first()
        assert refund is not None
        assert Decimal(str(refund.amount)) == Decimal("250.00")


def test_missing_damage_reason_rejected(client, app_instance, seed_data):
    """RULE-10, TC-11: Damaged return without reason code is blocked."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    _, txn2_id, rental2_id, _, _ = create_delivery_rental(
        client, app_instance, seed_data["product_id"], seed_data["customer_id"], "4.5"
    )

    client.post(
        f"/admin/transaction/{txn2_id}/return",
        data={
            "returned_items": str(rental2_id),
            f"qty_{rental2_id}": "1",
            f"status_{rental2_id}": "Full",
            f"condition_{rental2_id}": "Major Damage",
            "return_notes": "No reason selected",
            "late_fees": "0",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        rental2 = db.session.get(Rental, rental2_id)
        assert rental2.status == "Active"
        assert Decimal(str(rental2.damage_fee or 0)) == Decimal("0.00")


def test_rule_09_late_fee_calculation_and_audit(client, app_instance, seed_data):
    """RULE-09, TC-10: Late fee computed at rate * days, written to audit log."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    late_return_date = date.today() - timedelta(days=10)
    late_start_date = date.today() - timedelta(days=40)

    _, late_txn_id, late_rental_id, _, _ = create_delivery_rental(
        client,
        app_instance,
        seed_data["product_id"],
        seed_data["customer_id"],
        "2",
        quantity=1,
        start_date=late_start_date.isoformat(),
        return_date=late_return_date.isoformat(),
    )

    client.post(
        f"/admin/transaction/{late_txn_id}/return",
        data={
            "returned_items": str(late_rental_id),
            f"qty_{late_rental_id}": "1",
            f"status_{late_rental_id}": "Full",
            f"condition_{late_rental_id}": "Good",
            "return_notes": "Late return audit test",
            "late_fees": "0",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        late_rental = db.session.get(Rental, late_rental_id)
        late_txn = db.session.get(Transaction, late_txn_id)
        expected_days_late = (datetime.now().date() - late_rental.expected_return_date).days
        expected_late_fee = Decimal("100.00") * expected_days_late

        assert Decimal(str(late_rental.late_fees_incurred or 0)) == expected_late_fee

        late_log = (
            SecurityLog.query.filter_by(event_type="Late Fee Assessed")
            .order_by(SecurityLog.id.desc())
            .first()
        )
        assert late_log is not None
        assert f"Txn {late_txn.reference_no}" in late_log.description
        assert f"₱{expected_late_fee:,.2f}" in late_log.description

        late_payment = Payment.query.filter_by(
            transaction_id=late_txn_id, payment_type="Late Fee"
        ).first()
        assert late_payment is not None
        assert Decimal(str(late_payment.amount)) == expected_late_fee


def test_rule_08_partial_returns_and_visibility(client, app_instance, seed_data):
    """RULE-08, TC-09: Partial return preserves outstanding units and dashboard flag."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    _, partial_txn_id, partial_rental_id, _, _ = create_delivery_rental(
        client, app_instance, seed_data["product_id"], seed_data["customer_id"], "2", quantity=2
    )

    client.post(
        f"/admin/transaction/{partial_txn_id}/return",
        data={
            "returned_items": str(partial_rental_id),
            f"qty_{partial_rental_id}": "1",
            f"status_{partial_rental_id}": "Full",
            f"condition_{partial_rental_id}": "Good",
            "return_notes": "Partial return test",
            "late_fees": "0",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        partial_rental = db.session.get(Rental, partial_rental_id)
        assert partial_rental.status == "Partially Returned"
        assert partial_rental.quantity_returned == 1
        assert partial_rental.remaining_to_return == 1

    r = client.get("/admin/active-rentals?filter=partial_return")
    assert r.status_code == 200
    assert b"Test Wheelchair" in r.data
    assert b"outstanding" in r.data

    r = client.get("/admin/notifications/data")
    assert r.status_code == 200
    data = r.get_json() or {}
    partial_notice = next(
        (n for n in data.get("notifications", []) if n.get("id") == "partial_returns"),
        None,
    )
    assert partial_notice is not None
    assert "1 unit still outstanding" in (partial_notice.get("message") or "")

    # Close the remaining return
    client.post(
        f"/admin/transaction/{partial_txn_id}/return",
        data={
            "returned_items": str(partial_rental_id),
            f"qty_{partial_rental_id}": "1",
            f"status_{partial_rental_id}": "Full",
            f"condition_{partial_rental_id}": "Good",
            "return_notes": "Partial return completion test",
            "late_fees": "0",
            "payment_method": "Cash",
        },
        follow_redirects=True,
    )

    with app_instance.app_context():
        partial_rental = db.session.get(Rental, partial_rental_id)
        assert partial_rental.status == "Returned"
        assert partial_rental.remaining_to_return == 0


def test_fr08_restock_and_archive_movement_logs(client, app_instance, seed_data):
    """FR08: Inventory movement logs carry structured reference, user, and reason code."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    with app_instance.app_context():
        audit_product = Product(
            equipment_type="Wheelchair",
            name="Audit Restock Item",
            stock=5,
            reorder_level=2,
            cost_price=Decimal("100.00"),
            rent_price=Decimal("50.00"),
            sale_price=Decimal("200.00"),
            rent_period="Monthly",
            transaction_type="Both",
            condition="Brand New",
            is_refillable=False,
            status="Available",
            is_active=True,
        )
        db.session.add(audit_product)
        db.session.commit()
        audit_product_id = audit_product.id

    r = client.post(
        f"/admin/update_stock/{audit_product_id}",
        json={
            "increment": 3,
            "reason": "Supplier Invoice #8821",
            "total_amount_paid": 300,
        },
    )
    assert r.status_code == 200
    assert (r.get_json() or {}).get("success")

    with app_instance.app_context():
        restock_log = (
            InventoryLog.query.filter_by(product_id=audit_product_id, reason_code="RESTOCK")
            .order_by(InventoryLog.id.desc())
            .first()
        )
        assert restock_log is not None
        assert (restock_log.reference_no or "").startswith("RST-")
        assert restock_log.user_id == seed_data["admin_id"]
        assert restock_log.quantity == 3
        assert restock_log.balance_after == 8

    # Archive product
    r = client.post(f"/admin/delete-product/{audit_product_id}", follow_redirects=True)
    assert r.status_code == 200

    with app_instance.app_context():
        archive_log = (
            InventoryLog.query.filter_by(product_id=audit_product_id, action="ARCHIVE")
            .order_by(InventoryLog.id.desc())
            .first()
        )
        assert archive_log is not None
        assert (archive_log.reference_no or "").startswith("ARC-")
        assert archive_log.reason_code == "ARCHIVE"
        assert archive_log.quantity == -8
        assert archive_log.balance_after == 0


def test_fr06_primegas_refill_and_supplier_reconciliation(client, app_instance):
    """FR06: Primegas refill updates cylinder status and populates supplier reconciliation export."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    with app_instance.app_context():
        tank_product = Product(
            equipment_type="Oxygen Tank",
            name="Audit Refill Tank",
            stock=0,
            reorder_level=2,
            cost_price=Decimal("100.00"),
            rent_price=Decimal("200.00"),
            sale_price=Decimal("500.00"),
            rent_period="Monthly",
            transaction_type="Both",
            condition="Brand New",
            is_refillable=True,
            status="Available",
            is_active=True,
        )
        db.session.add(tank_product)
        db.session.flush()
        db.session.add(TankStatus(
            product_id=tank_product.id,
            total_owned=5,
            full_in_stock=0,
            empty_in_stock=5,
            rented_out=0,
            damaged_in_stock=0,
        ))
        db.session.commit()
        tank_product_id = tank_product.id

    r = client.post(
        "/admin/process-primegas",
        data={"product_id": str(tank_product_id), "quantity": "5", "amount": "500"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app_instance.app_context():
        tank_product = db.session.get(Product, tank_product_id)
        assert tank_product.tank_status.empty_in_stock == 0
        assert tank_product.tank_status.full_in_stock == 5

        refill_log = (
            InventoryLog.query.filter_by(product_id=tank_product_id, reason_code="REFILL")
            .order_by(InventoryLog.id.desc())
            .first()
        )
        assert refill_log is not None
        assert (refill_log.reference_no or "").startswith("EXP-")
        assert refill_log.quantity == 5
        assert refill_log.balance_after == 5

        primegas_expense = (
            Expense.query.filter_by(category="Primegas Refill", product_id=tank_product_id)
            .order_by(Expense.id.desc())
            .first()
        )
        assert primegas_expense is not None
        assert primegas_expense.refill_quantity == 5
        assert Decimal(str(primegas_expense.amount)) == Decimal("500.00")

    r = client.get("/admin/reports/export/supplier")
    assert r.status_code == 200
    assert "spreadsheet" in (r.headers.get("Content-Type") or "")

    wb = openpyxl.load_workbook(BytesIO(r.data))
    ws = wb["Prime Gas Reconciliation"]
    headers = [cell.value for cell in ws[4]]
    assert "Current Empty Awaiting Refill" in headers
    assert "Current Full in Stock" in headers

    supplier_rows = list(ws.iter_rows(min_row=5, values_only=True))
    tank_row = next((row for row in supplier_rows if row[2] == "Audit Refill Tank"), None)
    assert tank_row is not None
    assert tank_row[3] == 5
    assert tank_row[4] == 0
    assert tank_row[5] == 5


def test_tc04_atomic_stock_prevents_overselling(client, app_instance, seed_data):
    """TC-04: Atomic stock checking prevents racing transactions from taking stock beyond 0."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    with app_instance.app_context():
        race_product = Product(
            equipment_type="Wheelchair",
            name="Audit Race Item",
            stock=1,
            reorder_level=0,
            cost_price=Decimal("100.00"),
            rent_price=Decimal("100.00"),
            sale_price=Decimal("300.00"),
            rent_period="Monthly",
            transaction_type="Both",
            condition="Brand New",
            is_refillable=False,
            status="Available",
            is_active=True,
        )
        db.session.add(race_product)
        db.session.commit()
        race_product_id = race_product.id

    race_data = {
        "product_ids": str(race_product_id),
        f"quantity_{race_product_id}": "1",
        f"unit_price_{race_product_id}": "100",
        "start_date": "2026-10-01",
        "return_date": "2026-11-30",
        "duration_unit": "months",
        "customer_id": str(seed_data["customer_id"]),
        "fulfillment_type": "Delivery",
        "distance_km": "2",
        "delivery_address": "Jaro, Iloilo City",
        "delivery_fee": "100",
        "customer_deposit": "0",
        "amount_paid": "0",
        "payment_method": "Cash",
    }
    r1 = client.post("/admin/process-rental", data=race_data, follow_redirects=True)
    r2 = client.post("/admin/process-rental", data=race_data, follow_redirects=True)

    with app_instance.app_context():
        race_product = db.session.get(Product, race_product_id)
        race_rentals = Rental.query.filter_by(product_id=race_product_id).all()
        assert len(race_rentals) == 1
        assert (race_product.stock or 0) == 0

    assert b"Not enough stock" in r2.data


def test_audit_log_immutability_and_soft_archive(client, app_instance, seed_data):
    """NFR Reliability: Staff cannot archive logs; Admin soft-archives only with reason."""
    with app_instance.app_context():
        target_log = SecurityLog.query.filter_by(event_type="Damage Approved").first()
        target_log_id = target_log.id
        total_before = SecurityLog.query.count()

    logout(client)
    login(client, "test.staff@gmail.com", "Staff@12345")
    r = client.post(
        f"/admin/security/delete-log/{target_log_id}",
        data=json.dumps({"reason": "Staff attempted archive"}),
        content_type="application/json",
    )
    assert r.status_code == 403

    logout(client)
    login(client, "test.admin@gmail.com", "Admin@12345")

    # Missing reason rejected
    r = client.post(
        f"/admin/security/delete-log/{target_log_id}",
        data=json.dumps({}),
        content_type="application/json",
    )
    assert r.status_code == 400

    # Admin with reason succeeds
    r = client.post(
        f"/admin/security/delete-log/{target_log_id}",
        data=json.dumps({"reason": "Routine retention review"}),
        content_type="application/json",
    )
    assert r.status_code == 200
    assert r.get_json().get("success")

    with app_instance.app_context():
        archived = db.session.get(SecurityLog, target_log_id)
        assert archived is not None
        assert archived.is_archived
        assert archived.archive_reason == "Routine retention review"
        assert SecurityLog.query.filter_by(event_type="Audit Log Archived").count() >= 1

    # Bulk archive retains all rows
    r = client.post(
        "/admin/security/clear-logs",
        data={"reason": "Bulk archive test"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app_instance.app_context():
        total_after = SecurityLog.query.count()
        active_after = SecurityLog.query.filter_by(is_archived=False).count()
        assert total_after >= total_before
        assert active_after == 1
        assert SecurityLog.query.filter_by(event_type="Audit Logs Archived").count() == 1

    r = client.get("/admin/security?archived=1")
    assert r.status_code == 200


def test_settings_validation_and_tier_persistence(client, app_instance):
    """SystemSettings: Distance tiers validated and persisted correctly."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    r = client.get("/admin/settings")
    assert r.status_code == 200
    assert b"Distance-Based Deposit Tiers" in r.data

    # Invalid tier rejected
    r = client.post(
        "/admin/settings",
        data={
            "deposit_flat_amount": "1500",
            "deposit_distance_tiers": "not-a-tier",
            "late_fee_per_day": "100",
            "damage_fee_minor_pct": "25",
            "damage_fee_major_pct": "50",
            "default_reorder_level": "5",
        },
        follow_redirects=True,
    )
    assert b"not a tier" in r.data or b"Please check the values" in r.data

    # Valid tier saved
    r = client.post(
        "/admin/settings",
        data={
            "deposit_flat_amount": "1500",
            "deposit_distance_tiers": "2:400,5:900,*:1600",
            "late_fee_per_day": "100",
            "damage_fee_minor_pct": "25",
            "damage_fee_major_pct": "50",
            "default_reorder_level": "5",
        },
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app_instance.app_context():
        assert (
            SystemSetting.value_of("deposit_distance_tiers")
            == "2:400.00, 5:900.00, *:1600.00"
        )


def test_core_pages_render(client, app_instance):
    """Smoke test: dashboard, transactions, details, security logs, and delivery dispatch render with HTTP 200."""
    login(client, "test.admin@gmail.com", "Admin@12345")
    for path in ["/admin/dashboard", "/admin/transactions", "/admin/transaction_details/1", "/admin/security", "/admin/delivery-dispatch"]:
        r = client.get(path)
        assert r.status_code == 200


def test_delivery_dispatch_flow_and_status_updates(client, app_instance, seed_data):
    """2.1: Staff and Admin can view deliveries, filter by status, and update dispatch status."""
    login(client, "test.admin@gmail.com", "Admin@12345")

    # Create a rental delivery transaction
    _, txn_id, _, deposit, _ = create_delivery_rental(
        client, app_instance, seed_data["product_id"], seed_data["customer_id"], "4.5"
    )

    # 1. Delivery dispatch page renders and displays the new delivery
    r = client.get("/admin/delivery-dispatch")
    assert r.status_code == 200
    assert b"Delivery Dispatch" in r.data

    with app_instance.app_context():
        txn = db.session.get(Transaction, txn_id)
        assert txn.delivery_status == "Pending"

    # 2. Update status to 'Out for Delivery'
    r = client.post(
        f"/admin/delivery-dispatch/{txn_id}/status",
        data={"delivery_status": "Out for Delivery"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app_instance.app_context():
        txn = db.session.get(Transaction, txn_id)
        assert txn.delivery_status == "Out for Delivery"
        assert txn.tracking_status == "SHIPPED"

    # 3. Update status to 'Delivered'
    r = client.post(
        f"/admin/delivery-dispatch/{txn_id}/status",
        data={"delivery_status": "Delivered"},
        follow_redirects=True,
    )
    assert r.status_code == 200

    with app_instance.app_context():
        txn = db.session.get(Transaction, txn_id)
        assert txn.delivery_status == "Delivered"
        assert txn.tracking_status == "DELIVERED"

    # 4. Filter tabs work
    r_delivered = client.get("/admin/delivery-dispatch?status=Delivered")
    assert r_delivered.status_code == 200

    # 5. Staff with permission can view dispatch
    logout(client)
    login(client, "test.staff@gmail.com", "Staff@12345")
    r_staff = client.get("/admin/delivery-dispatch")
    assert r_staff.status_code == 200


def test_backup_restore_with_automatic_safety_backup(client, app_instance, seed_data):
    """2.2: Database restore creates an automatic pre-restore safety backup and recovers data."""
    import os
    from utils.backup import BACKUP_DIR, create_backup, restore_backup
    from models.customer import Customer

    login(client, "test.admin@gmail.com", "Admin@12345")

    # Step 1: Create a test backup
    success, backup_filename = create_backup(triggered_by="test_run")
    assert success, f"Failed to create test backup: {backup_filename}"
    backup_path = os.path.join(BACKUP_DIR, backup_filename)
    assert os.path.exists(backup_path)

    created_backups = [backup_path]

    try:
        # Step 2: Make a modification in the database
        with app_instance.app_context():
            cust = db.session.get(Customer, seed_data["customer_id"])
            original_last_name = cust.last_name
            cust.last_name = "ModifiedForRestoreTest"
            db.session.commit()

            cust_check = db.session.get(Customer, seed_data["customer_id"])
            assert cust_check.last_name == "ModifiedForRestoreTest"

        # Step 3: Restore the backup
        restore_ok, restore_msg = restore_backup(backup_filename, triggered_by="admin")
        assert restore_ok, f"Restore failed: {restore_msg}"

        # Step 4: Verify the pre-restore safety backup was automatically created
        safety_backups = [
            os.path.join(BACKUP_DIR, f)
            for f in os.listdir(BACKUP_DIR)
            if "_pre_restore_safety_" in f
        ]
        assert len(safety_backups) >= 1, "Expected pre-restore safety backup file to exist"
        created_backups.extend(safety_backups)

        # Step 5: Verify the database was restored to original
        with app_instance.app_context():
            cust_restored = db.session.get(Customer, seed_data["customer_id"])
            assert cust_restored is not None
            assert cust_restored.last_name == original_last_name

    finally:
        # Clean up temporary test backup files
        for b_path in created_backups:
            if os.path.exists(b_path):
                try:
                    os.remove(b_path)
                except Exception:
                    pass

